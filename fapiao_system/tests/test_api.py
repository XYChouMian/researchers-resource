"""API 冒烟测试：登录、记录增删改、发票/附件上传删除、选中导出 zip、报销/恢复/已处理、预览、日志、用户管理、迁移、子路径。"""
import io
import json
import zipfile

import pytest

import db
from PIL import Image

from conftest import ADMIN_ID, ADMIN_PASSWORD, make_user

PDF_BYTES = b"%PDF-1.4\n fake invoice"


def _png_bytes():
    buf = io.BytesIO()
    Image.new("RGB", (30, 60), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def login(client, student_id, password):
    return client.post("/api/login", json={"student_id": student_id, "password": password})


def create_record(client, **overrides):
    payload = {
        "product_name": "移液枪枪头",
        "channel": "taobao",
        "channel_note": "",
        "amount": "128.5",
        "paid_at": "2026-09-01",
        "payer": "self",
    }
    payload.update(overrides)
    return client.post("/api/records", json=payload)


def upload_files(client, rid, items=(("fapiao1.pdf", "invoice"), ("fapiao2.jpg", "invoice")), remarks=None, content=None):
    data = {
        "files": [(io.BytesIO(content if content is not None else PDF_BYTES), name) for name, _ in items],
        "categories": [category for _, category in items],
    }
    if remarks is not None:
        data["remarks"] = remarks
    return client.post(
        f"/api/records/{rid}/invoices",
        data=data,
        content_type="multipart/form-data",
    )


def test_login_wrong_password(client):
    assert login(client, "admin", "wrong").status_code == 401


def test_unauthorized_access(client):
    assert client.get("/api/records").status_code == 401
    assert client.get("/api/admin/users").status_code == 401


def test_user_cannot_access_admin_api(client):
    make_user(client)
    login(client, "2023123", "2023123")
    assert client.get("/api/admin/users").status_code == 403
    assert client.post("/api/admin/reimburse", json={"ids": [1]}).status_code == 403


def test_record_validation(client):
    make_user(client)
    login(client, "2023123", "2023123")
    assert create_record(client, channel="other", channel_note="").status_code == 400
    assert create_record(client, amount="abc").status_code == 400
    assert create_record(client, amount="0").status_code == 400
    assert create_record(client, paid_at="2026/09/01").status_code == 400
    assert create_record(client, payer="boss").status_code == 400
    assert client.get("/api/records").get_json()["records"] == []


def test_full_reimburse_flow(client):
    make_user(client)
    login(client, "2023123", "2023123")

    resp = create_record(client)
    assert resp.status_code == 201
    rid = resp.get_json()["record"]["id"]

    resp = client.put(f"/api/records/{rid}", json={
        "product_name": "移液枪枪头（改）", "channel": "jd", "channel_note": "",
        "amount": "99.9", "paid_at": "2026-09-02", "payer": "tang",
    })
    assert resp.status_code == 200
    assert resp.get_json()["record"]["channel"] == "jd"

    resp = upload_files(client, rid)
    assert resp.status_code == 200
    record = resp.get_json()["record"]
    assert record["status"] == "invoiced"
    assert record["invoice_count"] == 2
    assert len(record["invoices"]) == 2

    # 已开票的记录仍然可以修改
    resp = client.put(f"/api/records/{rid}", json={
        "product_name": "移液枪枪头（改2）", "channel": "jd", "channel_note": "",
        "amount": "88", "paid_at": "2026-09-02", "payer": "tang",
    })
    assert resp.status_code == 200

    # 已开票后可以补充附件，开票张数不受影响
    resp = upload_files(client, rid, (("hetong.pdf", "attachment"),))
    assert resp.status_code == 200
    record = resp.get_json()["record"]
    assert record["status"] == "invoiced"
    assert record["invoice_count"] == 2
    assert len(record["invoices"]) == 3

    inv_id = record["invoices"][0]["id"]
    resp = client.get(f"/api/invoices/{inv_id}/download")
    assert resp.status_code == 200
    assert resp.data.startswith(b"%PDF")

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    resp = client.get(f"/api/admin/export?ids={rid}")
    assert resp.status_code == 200
    assert resp.mimetype == "application/zip"
    assert resp.data[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(resp.data)) as zf:
        assert "报销清单.xlsx" in zf.namelist()

    resp = client.post("/api/admin/reimburse", json={"ids": [rid]})
    assert resp.status_code == 200
    assert resp.get_json()["updated"] == 1

    login(client, "2023123", "2023123")
    record = client.get("/api/records").get_json()["records"][0]
    assert record["status"] == "reimbursed"
    # 已报销后：上传、改、删记录、删文件全部被锁
    assert upload_files(client, rid, (("x.pdf", "invoice"),)).status_code == 400
    assert client.put(f"/api/records/{rid}", json={
        "product_name": "x", "channel": "jd", "channel_note": "",
        "amount": "1", "paid_at": "2026-09-02", "payer": "tang",
    }).status_code == 400
    assert client.delete(f"/api/records/{rid}").status_code == 400
    assert client.delete(f"/api/invoices/{inv_id}").status_code == 400


def test_invoice_count_auto_and_revert(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]

    resp = upload_files(client, rid, (
        ("f1.pdf", "invoice"), ("a1.jpg", "attachment"), ("f2.png", "invoice"),
    ))
    record = resp.get_json()["record"]
    assert record["status"] == "invoiced"
    assert record["invoice_count"] == 2
    assert [i["category"] for i in record["invoices"]] == [
        "invoice", "attachment", "invoice",
    ]

    resp = client.delete(f"/api/invoices/{record['invoices'][0]['id']}")
    record = resp.get_json()["record"]
    assert record["invoice_count"] == 1
    assert record["status"] == "invoiced"

    att_id = next(i["id"] for i in record["invoices"] if i["category"] == "attachment")
    record = client.delete(f"/api/invoices/{att_id}").get_json()["record"]
    assert record["invoice_count"] == 1
    assert record["status"] == "invoiced"

    last_inv_id = next(i["id"] for i in record["invoices"] if i["category"] == "invoice")
    record = client.delete(f"/api/invoices/{last_inv_id}").get_json()["record"]
    assert record["invoice_count"] == 0
    assert record["status"] == "pending"


def test_upload_categories_validation(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]

    resp = client.post(
        f"/api/records/{rid}/invoices",
        data={
            "files": [(io.BytesIO(PDF_BYTES), "a.pdf")],
            "categories": ["invoice", "attachment"],
        },
        content_type="multipart/form-data",
    )
    assert resp.status_code == 400

    resp = client.post(
        f"/api/records/{rid}/invoices",
        data={
            "files": [(io.BytesIO(PDF_BYTES), "a.pdf")],
            "categories": ["bogus"],
        },
        content_type="multipart/form-data",
    )
    assert resp.status_code == 400

    resp = client.post(
        f"/api/records/{rid}/invoices",
        data={"files": [(io.BytesIO(PDF_BYTES), "a.pdf")]},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 400
    assert client.get("/api/records").get_json()["records"][0]["invoices"] == []


def test_upload_rejects_bad_extension(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    resp = client.post(
        f"/api/records/{rid}/invoices",
        data={
            "files": [(io.BytesIO(b"MZxxx"), "virus.exe")],
            "categories": ["invoice"],
        },
        content_type="multipart/form-data",
    )
    assert resp.status_code == 400
    assert client.get("/api/records").get_json()["records"][0]["status"] == "pending"


def test_delete_record_allowed_until_reimbursed(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("a.pdf", "invoice"),))
    rid2 = create_record(client).get_json()["record"]["id"]
    assert client.delete(f"/api/records/{rid2}").status_code == 200
    assert client.delete(f"/api/records/{rid}").status_code == 200
    assert client.get("/api/records").get_json()["records"] == []


def test_user_cannot_touch_others_record(client):
    make_user(client, "2023123")
    make_user(client, "2023456")
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("a.pdf", "invoice"),))
    inv_id = client.get("/api/records").get_json()["records"][0]["invoices"][0]["id"]
    login(client, "2023456", "2023456")
    assert client.delete(f"/api/records/{rid}").status_code == 404
    assert upload_files(client, rid, (("x.pdf", "invoice"),)).status_code == 404
    assert client.get(f"/api/invoices/{inv_id}/download").status_code == 403
    assert client.delete(f"/api/invoices/{inv_id}").status_code == 403


def test_change_password(client):
    make_user(client)
    login(client, "2023123", "2023123")
    resp = client.post("/api/password", json={"old_password": "bad", "new_password": "654321"})
    assert resp.status_code == 400
    resp = client.post("/api/password", json={"old_password": "2023123", "new_password": "654321"})
    assert resp.status_code == 200
    assert login(client, "2023123", "2023123").status_code == 401
    assert login(client, "2023123", "654321").status_code == 200


def test_admin_user_management(client):
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/users", json={"student_id": "2023123"}).status_code == 201
    assert client.post("/api/admin/users", json={"student_id": "2023123"}).status_code == 400
    assert client.post("/api/admin/users", json={"student_id": "非法!"}).status_code == 400

    users = client.get("/api/admin/users").get_json()["users"]
    uid = next(u["id"] for u in users if u["student_id"] == "2023123")
    admin_id = next(u["id"] for u in users if u["student_id"] == ADMIN_ID)

    assert client.post(f"/api/admin/users/{uid}", json={}).status_code == 405
    assert client.post(f"/api/admin/users/{uid}/reset_password", json={}).status_code == 200
    assert client.delete(f"/api/admin/users/{admin_id}").status_code == 400

    resp = client.put(f"/api/admin/users/{uid}", json={"name": "新名字", "role": "user"})
    assert resp.status_code == 200

    resp = client.post("/api/admin/users", json={"student_id": "2099999"})
    uid2 = resp.get_json()["user"]["id"]
    assert client.delete(f"/api/admin/users/{uid2}").status_code == 200


def test_admin_cannot_delete_user_with_records(client):
    make_user(client)
    login(client, "2023123", "2023123")
    create_record(client)
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    users = client.get("/api/admin/users").get_json()["users"]
    uid = next(u["id"] for u in users if u["student_id"] == "2023123")
    resp = client.delete(f"/api/admin/users/{uid}")
    assert resp.status_code == 400
    assert "购买记录" in resp.get_json()["error"]


def test_upload_remark(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]

    record = upload_files(client, rid, (("f.pdf", "invoice"),), remarks=["增值税发票"]).get_json()["record"]
    assert record["invoices"][0]["remark"] == "增值税发票"

    record = upload_files(client, rid, (("a.jpg", "attachment"),)).get_json()["record"]
    assert record["invoices"][1]["remark"] == ""

    resp = upload_files(client, rid, (("b.jpg", "attachment"),), remarks=["x" * 101])
    assert resp.status_code == 400
    resp = upload_files(client, rid, (("c.jpg", "attachment"),), remarks=["a", "b"])
    assert resp.status_code == 400


def test_invoice_inline_preview(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("f.pdf", "invoice"),))
    inv_id = client.get("/api/records").get_json()["records"][0]["invoices"][0]["id"]

    resp = client.get(f"/api/invoices/{inv_id}/download")
    assert "attachment" in resp.headers.get("Content-Disposition", "")
    assert resp.mimetype == "application/pdf"

    resp = client.get(f"/api/invoices/{inv_id}/download?inline=1")
    assert "attachment" not in resp.headers.get("Content-Disposition", "")
    assert resp.mimetype == "application/pdf"
    assert resp.data.startswith(b"%PDF")


def test_image_preview_converted_to_pdf(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    resp = upload_files(client, rid, (("photo.png", "invoice"),), content=_png_bytes())
    inv_id = resp.get_json()["record"]["invoices"][0]["id"]
    png = _png_bytes()

    resp = client.get(f"/api/invoices/{inv_id}/download")
    assert resp.mimetype == "image/png"
    assert resp.data == png

    resp = client.get(f"/api/invoices/{inv_id}/download?inline=1")
    assert resp.status_code == 200
    assert resp.mimetype == "application/pdf"
    assert resp.data.startswith(b"%PDF")

    folder = db.UPLOAD_DIR / "2023123"
    conn = db.connect()
    stored_name = conn.execute("SELECT stored_name FROM invoices").fetchone()[0]
    conn.close()
    original = folder / stored_name
    derived = folder / (stored_name.rsplit(".", 1)[0] + ".preview.pdf")
    assert original.is_file()
    assert derived.is_file()

    client.delete(f"/api/invoices/{inv_id}")
    assert not original.exists()
    assert not derived.exists()


def test_corrupt_image_fallback(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    resp = upload_files(client, rid, (("broken.jpg", "invoice"),))
    inv_id = resp.get_json()["record"]["invoices"][0]["id"]

    resp = client.get(f"/api/invoices/{inv_id}/download?inline=1")
    assert resp.status_code == 200
    assert resp.mimetype == "image/jpeg"
    assert resp.data == PDF_BYTES


def test_admin_records_filter(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.get("/api/admin/records?status=invoiced").get_json()["records"] == []
    assert client.get("/api/admin/records?status=processed").get_json()["records"] == []
    data = client.get("/api/admin/records").get_json()
    assert len(data["records"]) == 1
    assert data["records"][0]["student_id"] == "2023123"

    client.post("/api/admin/process", json={"ids": [rid], "note": "x"})
    for status in ("pending", "invoiced", "reimbursed"):
        assert client.get(f"/api/admin/records?status={status}").get_json()["records"] == []
    processed = client.get("/api/admin/records?status=processed").get_json()["records"]
    assert [r["status"] for r in processed] == ["processed"]


def test_reimburse_invalid_ids(client):
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/reimburse", json={"ids": []}).status_code == 400
    assert client.post("/api/admin/reimburse", json={"ids": ["abc"]}).status_code == 400


def test_reimburse_requires_invoiced(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    resp = client.post("/api/admin/reimburse", json={"ids": [rid]})
    assert resp.status_code == 200
    assert resp.get_json()["updated"] == 0
    assert client.get("/api/admin/records?status=pending").get_json()["records"][0]["status"] == "pending"

    login(client, "2023123", "2023123")
    upload_files(client, rid, (("f.pdf", "invoice"),))
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/reimburse", json={"ids": [rid]}).get_json()["updated"] == 1

    assert client.post("/api/admin/revert", json={"ids": [rid]}).get_json()["updated"] == 1
    record = client.get("/api/admin/records?status=invoiced").get_json()["records"][0]
    assert record["status"] == "invoiced"

    assert client.post("/api/admin/reimburse", json={"ids": [rid]}).get_json()["updated"] == 1


def test_process_and_revert(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/process", json={"ids": [rid], "note": "   "}).status_code == 400
    assert client.post("/api/admin/process", json={"ids": [rid], "note": "x" * 201}).status_code == 400
    assert client.post("/api/admin/process", json={"ids": [rid], "note": "线下退款处理"}).get_json()["updated"] == 1

    record = client.get("/api/admin/records?status=processed").get_json()["records"][0]
    assert record["status"] == "processed"
    assert record["process_note"] == "线下退款处理"

    login(client, "2023123", "2023123")
    assert upload_files(client, rid, (("x.pdf", "invoice"),)).status_code == 400
    assert client.put(f"/api/records/{rid}", json={
        "product_name": "x", "channel": "jd", "channel_note": "",
        "amount": "1", "paid_at": "2026-09-02", "payer": "tang",
    }).status_code == 400
    assert client.delete(f"/api/records/{rid}").status_code == 400

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/revert", json={"ids": [rid]}).get_json()["updated"] == 1
    record = client.get("/api/admin/records?status=pending").get_json()["records"][0]
    assert record["status"] == "pending"
    assert record["process_note"] is None

    login(client, "2023123", "2023123")
    upload_files(client, rid, (("f.pdf", "invoice"),))
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/process", json={"ids": [rid], "note": "发票作废重开"}).get_json()["updated"] == 1
    assert client.post("/api/admin/revert", json={"ids": [rid]}).get_json()["updated"] == 1
    record = client.get("/api/admin/records?status=invoiced").get_json()["records"][0]
    assert record["status"] == "invoiced"
    assert client.post("/api/admin/reimburse", json={"ids": [rid]}).get_json()["updated"] == 1

    login(client, "2023123", "2023123")
    rid2 = create_record(client).get_json()["record"]["id"]
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/revert", json={"ids": [rid2]}).get_json()["updated"] == 0


def test_reject_and_user_fix_flow(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("f.pdf", "invoice"),))
    rid2 = create_record(client).get_json()["record"]["id"]

    def user_record(rid_):
        records = client.get("/api/records").get_json()["records"]
        return next(r for r in records if r["id"] == rid_)

    def admin_record(rid_):
        records = client.get("/api/admin/records").get_json()["records"]
        return next(r for r in records if r["id"] == rid_)

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/reject", json={"ids": [rid], "note": "  "}).status_code == 400
    assert client.post("/api/admin/reject", json={"ids": [rid], "note": "x" * 201}).status_code == 400
    assert client.post("/api/admin/reject", json={"ids": [rid2], "note": "理由"}).get_json()["updated"] == 0
    assert client.post("/api/admin/reject", json={"ids": [rid], "note": "发票抬头有误"}).get_json()["updated"] == 1

    record = admin_record(rid)
    assert record["status"] == "rejected"
    assert record["reject_note"] == "发票抬头有误"

    login(client, "2023123", "2023123")
    assert user_record(rid)["reject_note"] == "发票抬头有误"

    resp = client.put(f"/api/records/{rid}", json={
        "product_name": "移液枪枪头（改）", "channel": "jd", "channel_note": "",
        "amount": "99.9", "paid_at": "2026-09-02", "payer": "tang",
    })
    record = resp.get_json()["record"]
    assert record["status"] == "invoiced"
    assert record["reject_note"] is None

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    client.post("/api/admin/reject", json={"ids": [rid], "note": "附件不清晰"})
    login(client, "2023123", "2023123")
    record = upload_files(client, rid, (("a.jpg", "attachment"),)).get_json()["record"]
    assert record["status"] == "invoiced"
    assert record["reject_note"] is None

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    client.post("/api/admin/reject", json={"ids": [rid], "note": "全部重传"})
    login(client, "2023123", "2023123")
    inv_id = user_record(rid)["invoices"][0]["id"]
    record = client.delete(f"/api/invoices/{inv_id}").get_json()["record"]
    assert record["status"] == "pending"
    assert record["reject_note"] is None

    login(client, "2023123", "2023123")
    upload_files(client, rid, (("f2.pdf", "invoice"),))
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    client.post("/api/admin/reject", json={"ids": [rid], "note": "删除重报"})
    login(client, "2023123", "2023123")
    assert client.delete(f"/api/records/{rid}").status_code == 200

    login(client, "2023123", "2023123")
    upload_files(client, rid2, (("g.pdf", "invoice"),))
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    client.post("/api/admin/reject", json={"ids": [rid2], "note": "误操作驳回"})
    assert client.post("/api/admin/revert", json={"ids": [rid2]}).get_json()["updated"] == 1
    record = admin_record(rid2)
    assert record["status"] == "invoiced"
    assert record["reject_note"] is None


def test_admin_can_submit_records(client):
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    resp = create_record(client)
    assert resp.status_code == 201
    rid = resp.get_json()["record"]["id"]
    assert upload_files(client, rid, (("a.pdf", "invoice"),)).status_code == 200
    assert client.get("/api/records").get_json()["records"][0]["status"] == "invoiced"


def test_admin_logs_api(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("f.pdf", "invoice"),))
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    client.post("/api/admin/reimburse", json={"ids": [rid], "note": "日志接口测试"})

    resp = client.get("/api/admin/logs")
    assert resp.status_code == 200
    entries = resp.get_json()["entries"]
    assert entries[0]["action_text"] == "标记已报销"
    assert entries[0]["product_name"] == "移液枪枪头"
    assert entries[0]["note"] == "日志接口测试"

    login(client, "2023123", "2023123")
    assert client.get("/api/admin/logs").status_code == 403
    client.post("/api/logout")
    assert client.get("/api/admin/logs").status_code == 401

    log_path = db.DATA_DIR / "admin.log"
    with log_path.open("a", encoding="utf-8") as fh:
        fh.write("legacy 非 JSON 行\n")
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    entries = client.get("/api/admin/logs").get_json()["entries"]
    assert entries[0] == {"raw": "legacy 非 JSON 行"}
    assert entries[1]["action"] == "reimburse"


def test_user_batch_delete_and_export(client):
    make_user(client)
    login(client, "2023123", "2023123")
    r_invoiced = create_record(client, product_name="A").get_json()["record"]["id"]
    upload_files(client, r_invoiced, (("f.pdf", "invoice"),))
    r_pending = create_record(client, product_name="B").get_json()["record"]["id"]
    r_locked = create_record(client, product_name="C").get_json()["record"]["id"]
    r_rejected = create_record(client, product_name="D").get_json()["record"]["id"]
    upload_files(client, r_locked, (("g.pdf", "invoice"),))
    upload_files(client, r_rejected, (("h.pdf", "invoice"),))
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    client.post("/api/admin/reimburse", json={"ids": [r_locked]})
    client.post("/api/admin/reject", json={"ids": [r_rejected], "note": "重传"})

    login(client, "2023123", "2023123")
    resp = client.post("/api/records/batch-delete", json={
        "ids": [r_invoiced, r_pending, r_rejected, r_locked],
    })
    assert resp.get_json()["deleted"] == 3
    remaining = {r["id"]: r["status"] for r in client.get("/api/records").get_json()["records"]}
    assert remaining == {r_locked: "reimbursed"}

    make_user(client, "2023456")
    login(client, "2023456", "2023456")
    other_id = create_record(client, product_name="别人的").get_json()["record"]["id"]
    login(client, "2023123", "2023123")
    resp = client.post("/api/records/batch-delete", json={"ids": [other_id]})
    assert resp.get_json()["deleted"] == 0

    assert client.get("/api/records/export").status_code == 400
    assert client.get("/api/records/export?ids=abc").status_code == 400
    assert client.get("/api/records/export?ids=999999").status_code == 404

    resp = client.get(f"/api/records/export?ids={r_locked},{other_id}")
    assert resp.status_code == 200
    assert resp.data[:2] == b"PK"
    entries = _zip_entries(resp.data)
    assert "报销清单.xlsx" in entries
    rows = _sheet_rows(entries["报销清单.xlsx"])
    assert len(rows) == 1
    assert rows[0][0] == "2023123"
    assert rows[0][2] == "C"


def test_audit_log_and_detail(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("f.pdf", "invoice"),))

    detail = client.get(f"/api/records/{rid}/detail").get_json()
    assert detail["record"]["status"] == "invoiced"
    assert len(detail["record"]["invoices"]) == 1
    assert detail["history"] == []

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    resp = client.post("/api/admin/reimburse", json={"ids": [rid], "note": "x" * 201})
    assert resp.status_code == 400
    resp = client.post("/api/admin/reimburse", json={"ids": [rid], "note": "随9月材料提交"})
    assert resp.get_json()["updated"] == 1

    detail = client.get(f"/api/records/{rid}/detail").get_json()
    assert detail["record"]["reimburse_note"] == "随9月材料提交"
    assert [h["action"] for h in detail["history"]] == ["reimburse"]
    assert detail["history"][0]["operator"] == "管理员"

    log_text = (db.DATA_DIR / "admin.log").read_text(encoding="utf-8")
    entry = json.loads([line for line in log_text.splitlines() if line.strip()][-1])
    assert entry["action"] == "reimburse"
    assert entry["action_text"] == "标记已报销"
    assert entry["operator"] == "管理员"
    assert entry["operator_id"] == ADMIN_ID
    assert entry["record_id"] == rid
    assert entry["student_id"] == "2023123"
    assert entry["name"] == "测试同学"
    assert entry["product_name"] == "移液枪枪头"
    assert entry["amount"] == 128.5
    assert entry["paid_at"] == "2026-09-01"
    assert entry["channel"] == "淘宝"
    assert entry["payer"] == "本人"
    assert entry["from_status"] == "已开票"
    assert entry["to_status"] == "已报销"
    assert entry["note"] == "随9月材料提交"

    make_user(client, "2023456")
    login(client, "2023456", "2023456")
    assert client.get(f"/api/records/{rid}/detail").status_code == 403

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    client.post("/api/admin/revert", json={"ids": [rid]})
    login(client, "2023123", "2023123")
    detail = client.get(f"/api/records/{rid}/detail").get_json()
    assert detail["record"]["reimburse_note"] is None
    assert [h["action"] for h in detail["history"]] == ["reimburse", "revert"]

    client.delete(f"/api/records/{rid}")
    conn = db.connect()
    count = conn.execute(
        "SELECT COUNT(*) FROM audit_logs WHERE record_id=?", (rid,)
    ).fetchone()[0]
    conn.close()
    assert count == 2


OLD_RECORDS_SQL = """
CREATE TABLE records (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id       INTEGER NOT NULL REFERENCES users(id),
    product_name  TEXT    NOT NULL,
    channel       TEXT    NOT NULL CHECK (channel IN ('taobao', 'jd', 'other')),
    channel_note  TEXT    NOT NULL DEFAULT '',
    amount        REAL    NOT NULL CHECK (amount > 0),
    paid_at       TEXT    NOT NULL,
    payer         TEXT    NOT NULL CHECK (payer IN ('self', 'tang')),
    invoice_count INTEGER,
    status        TEXT    NOT NULL DEFAULT 'pending'
                  CHECK (status IN ('pending', 'invoiced', 'reimbursed')),
    reimbursed_at TEXT,
    created_at    TEXT    NOT NULL DEFAULT (datetime('now', 'localtime')),
    updated_at    TEXT    NOT NULL DEFAULT (datetime('now', 'localtime'))
);
CREATE INDEX idx_records_user ON records(user_id);
"""


def test_migrate_old_records_table(client):
    """带旧 CHECK 约束的库应被 init_db 自动重建，老数据保留且支持 processed。"""
    conn = db.connect()
    conn.executescript("DROP TABLE records;" + OLD_RECORDS_SQL)
    conn.execute(
        "INSERT INTO records (user_id, product_name, channel, amount, paid_at, payer)"
        " VALUES (1, '旧记录', 'taobao', 1, '2026-01-01', 'self')"
    )
    conn.commit()
    conn.close()

    db.init_db()

    conn = db.connect()
    conn.execute(
        "INSERT INTO records (user_id, product_name, channel, amount, paid_at, payer,"
        " status, process_note, reject_note, reimburse_note)"
        " VALUES (1, '新记录', 'jd', 2, '2026-01-02', 'self', 'rejected', '线下处理', '驳回理由', '报销备注')"
    )
    conn.commit()
    assert conn.execute("SELECT status FROM records WHERE product_name='旧记录'").fetchone()[0] == "pending"
    row = conn.execute(
        "SELECT process_note, reject_note, reimburse_note FROM records WHERE product_name='新记录'"
    ).fetchone()
    assert row[0] == "线下处理"
    assert row[1] == "驳回理由"
    assert row[2] == "报销备注"
    conn.close()


def _export(client, ids):
    return client.get(f"/api/admin/export?ids={','.join(map(str, ids))}")


def _zip_entries(data):
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _sheet_rows(xlsx_bytes):
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(xlsx_bytes)).active
    return [list(row) for row in ws.iter_rows(min_row=2, values_only=True)]


def test_export_ids_validation(client):
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.get("/api/admin/export").status_code == 400
    assert client.get("/api/admin/export?ids=abc").status_code == 400
    assert client.get("/api/admin/export?ids=999").status_code == 404


def test_export_selected_zip(client):
    make_user(client, "2023123", "张三")
    make_user(client, "2023456", "张三")
    make_user(client, "2023789", "李四")

    login(client, "2023123", "2023123")
    ra = create_record(client, remark="实验室共用耗材").get_json()["record"]["id"]
    upload_files(client, ra, (("f1.pdf", "invoice"), ("a1.jpg", "attachment")))
    rb = create_record(client).get_json()["record"]["id"]
    upload_files(client, rb, (("f2.pdf", "invoice"),))

    login(client, "2023456", "2023456")
    rc = create_record(client, product_name="离心管", amount="9.9").get_json()["record"]["id"]
    upload_files(client, rc, (("f3.pdf", "invoice"),))

    login(client, "2023789", "2023789")
    rd = create_record(client, product_name="耗材", amount="5").get_json()["record"]["id"]

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    resp = _export(client, [ra, rb, rc, rd])
    assert resp.status_code == 200
    entries = _zip_entries(resp.data)
    assert set(entries) == {
        "报销清单.xlsx",
        "张三/移液枪枪头_128.50_发票.pdf",
        "张三/移液枪枪头_128.50_附件.jpg",
        "张三/移液枪枪头_128.50_发票(1).pdf",
        "张三(2)/离心管_9.90_发票.pdf",
    }

    rows = _sheet_rows(entries["报销清单.xlsx"])
    assert len(rows) == 4
    assert rows[0][5] == "实验室共用耗材"
    assert not rows[1][5]
    assert rows[0][-1] == "移液枪枪头_128.50_发票.pdf;移液枪枪头_128.50_附件.jpg"
    assert rows[1][-1] == "移液枪枪头_128.50_发票(1).pdf"
    assert rows[2][-1] == "离心管_9.90_发票.pdf"
    assert rows[2][0] == "2023456"
    assert rows[3][10] == "待开票"
    assert rows[3][-1] == "未开票"


def test_record_remark(client):
    make_user(client)
    login(client, "2023123", "2023123")
    resp = create_record(client, remark="  实验室共用耗材  ")
    record = resp.get_json()["record"]
    assert record["remark"] == "实验室共用耗材"
    rid = record["id"]

    resp = client.put(f"/api/records/{rid}", json={
        "product_name": "移液枪枪头", "channel": "taobao", "channel_note": "",
        "remark": "改后的备注", "amount": "128.5", "paid_at": "2026-09-01", "payer": "self",
    })
    assert resp.get_json()["record"]["remark"] == "改后的备注"

    assert create_record(client, remark="x" * 201).status_code == 400
    assert create_record(client).get_json()["record"]["remark"] == ""


def test_export_marks_missing_file(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("f.pdf", "invoice"),))
    conn = db.connect()
    stored = conn.execute("SELECT stored_name FROM invoices").fetchone()[0]
    conn.close()
    (db.UPLOAD_DIR / "2023123" / stored).unlink()

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    entries = _zip_entries(_export(client, [rid]).data)
    assert set(entries) == {"报销清单.xlsx"}
    rows = _sheet_rows(entries["报销清单.xlsx"])
    assert rows[0][-1] == "（文件缺失）"


def test_prefix_routing(client, monkeypatch):
    monkeypatch.setenv("FAPIAO_PREFIX", "/fapiao")
    from app import create_app

    app = create_app()
    app.config["TESTING"] = True
    c = app.test_client()
    assert c.get("/fapiao/api/me").status_code == 401
    assert c.get("/fapiao/").status_code == 200
    assert c.get("/fapiao/static/css/style.css").status_code == 200
    assert c.get("/api/me").status_code == 404
    assert c.get("/other/api/me").status_code == 404


PROTECTED_ENDPOINTS = [
    ("get", "/api/me", None),
    ("post", "/api/password", {"old_password": "x", "new_password": "yyyyyy"}),
    ("get", "/api/records", None),
    ("post", "/api/records", {}),
    ("put", "/api/records/1", {}),
    ("delete", "/api/records/1", None),
    ("post", "/api/records/1/invoices", None),
    ("delete", "/api/invoices/1", None),
    ("get", "/api/records/1/detail", None),
    ("get", "/api/invoices/1/download", None),
    ("post", "/api/records/batch-delete", {}),
    ("get", "/api/records/export", None),
    ("get", "/api/admin/users", None),
    ("post", "/api/admin/users", {}),
    ("put", "/api/admin/users/1", {}),
    ("post", "/api/admin/users/1/reset_password", {}),
    ("delete", "/api/admin/users/1", None),
    ("get", "/api/admin/records", None),
    ("get", "/api/admin/export", None),
    ("post", "/api/admin/reimburse", {}),
    ("post", "/api/admin/reject", {}),
    ("post", "/api/admin/revert", {}),
    ("post", "/api/admin/process", {}),
    ("get", "/api/admin/logs", None),
]

ADMIN_ENDPOINTS = PROTECTED_ENDPOINTS[12:]


def _call(client, method, path, payload):
    if payload is None:
        return getattr(client, method)(path)
    return getattr(client, method)(path, json=payload)


@pytest.mark.parametrize("method,path,payload", PROTECTED_ENDPOINTS)
def test_unauthorized_endpoints(client, method, path, payload):
    assert _call(client, method, path, payload).status_code == 401


@pytest.mark.parametrize("method,path,payload", ADMIN_ENDPOINTS)
def test_user_forbidden_admin_endpoints(client, method, path, payload):
    make_user(client)
    login(client, "2023123", "2023123")
    assert _call(client, method, path, payload).status_code == 403


def test_put_others_record_forbidden(client):
    make_user(client, "2023123")
    make_user(client, "2023456")
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]

    login(client, "2023456", "2023456")
    resp = client.put(f"/api/records/{rid}", json={
        "product_name": "改", "channel": "jd", "channel_note": "",
        "remark": "", "amount": "1", "paid_at": "2026-09-01", "payer": "self",
    })
    assert resp.status_code == 404
    assert client.get("/api/records").get_json()["records"] == []


def test_password_stored_hashed(client):
    make_user(client, "2023123")
    conn = db.connect()
    stored = conn.execute(
        "SELECT password_hash FROM users WHERE student_id='2023123'"
    ).fetchone()[0]
    conn.close()
    assert stored != "2023123"
    assert stored.startswith(("pbkdf2:", "scrypt:"))


def test_reset_password_effect(client):
    make_user(client)
    login(client, "2023123", "2023123")
    assert client.post("/api/password", json={
        "old_password": "2023123", "new_password": "newpass1",
    }).status_code == 200
    assert login(client, "2023123", "2023123").status_code == 401
    assert login(client, "2023123", "newpass1").status_code == 200

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    users = client.get("/api/admin/users").get_json()["users"]
    uid = next(u["id"] for u in users if u["student_id"] == "2023123")
    assert client.post(f"/api/admin/users/{uid}/reset_password", json={}).status_code == 200
    assert login(client, "2023123", "newpass1").status_code == 401
    assert login(client, "2023123", "2023123").status_code == 200


def test_change_password_too_short(client):
    make_user(client)
    login(client, "2023123", "2023123")
    resp = client.post("/api/password", json={
        "old_password": "2023123", "new_password": "12345",
    })
    assert resp.status_code == 400
    assert "6 位" in resp.get_json()["error"]


def test_record_field_boundaries(client):
    make_user(client)
    login(client, "2023123", "2023123")
    assert create_record(client, amount="1000000").status_code == 201
    assert create_record(client, amount="1000000.01").status_code == 400
    assert create_record(client, product_name="x" * 201).status_code == 400
    assert create_record(client, channel="other", channel_note="y" * 101).status_code == 400
    assert create_record(client, channel="other", channel_note="y" * 100).status_code == 201


def test_upload_empty_file(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    resp = client.post(
        f"/api/records/{rid}/invoices",
        data={"files": [(io.BytesIO(b""), "empty.pdf")], "categories": ["invoice"]},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 400
    assert "内容为空" in resp.get_json()["error"]


def test_upload_oversize_file(client, monkeypatch):
    import records
    monkeypatch.setattr(records, "MAX_INVOICE_SIZE", 10)
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    resp = upload_files(client, rid, (("big.pdf", "invoice"),))
    assert resp.status_code == 400
    assert "超过" in resp.get_json()["error"]


def test_request_too_large_413(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    client.application.config["MAX_CONTENT_LENGTH"] = 100
    resp = client.post(
        f"/api/records/{rid}/invoices",
        data={"files": [(io.BytesIO(PDF_BYTES), "a.pdf")], "categories": ["invoice"]},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 413
    assert "过大" in resp.get_json()["error"]


def test_download_missing_file_404(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("a.pdf", "invoice"),))
    inv_id = client.get("/api/records").get_json()["records"][0]["invoices"][0]["id"]
    conn = db.connect()
    stored = conn.execute("SELECT stored_name FROM invoices").fetchone()[0]
    conn.close()
    (db.UPLOAD_DIR / "2023123" / stored).unlink()

    resp = client.get(f"/api/invoices/{inv_id}/download")
    assert resp.status_code == 404
    assert "已丢失" in resp.get_json()["error"]


def test_admin_can_download_others_invoice(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("a.pdf", "invoice"),))
    inv_id = client.get("/api/records").get_json()["records"][0]["invoices"][0]["id"]

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.get(f"/api/invoices/{inv_id}/download").status_code == 200
    assert client.get(f"/api/invoices/{inv_id}/download?inline=1").status_code == 200


def test_delete_record_cleans_files(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid = create_record(client).get_json()["record"]["id"]
    upload_files(client, rid, (("photo.png", "invoice"),), content=_png_bytes())
    inv_id = client.get("/api/records").get_json()["records"][0]["invoices"][0]["id"]
    assert client.get(f"/api/invoices/{inv_id}/download?inline=1").status_code == 200

    folder = db.UPLOAD_DIR / "2023123"
    assert len(list(folder.iterdir())) == 2
    assert client.delete(f"/api/records/{rid}").status_code == 200
    assert list(folder.iterdir()) == []


def test_batch_delete_cleans_files(client):
    make_user(client)
    login(client, "2023123", "2023123")
    rid1 = create_record(client, product_name="A").get_json()["record"]["id"]
    upload_files(client, rid1, (("a.pdf", "invoice"),))
    rid2 = create_record(client, product_name="B").get_json()["record"]["id"]
    upload_files(client, rid2, (("b.jpg", "invoice"),))

    folder = db.UPLOAD_DIR / "2023123"
    assert len(list(folder.iterdir())) == 2
    resp = client.post("/api/records/batch-delete", json={"ids": [rid1, rid2]})
    assert resp.get_json()["deleted"] == 2
    assert list(folder.iterdir()) == []


def test_admin_user_validation(client):
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    assert client.post("/api/admin/users", json={
        "student_id": "2023123", "role": "boss",
    }).status_code == 400
    assert client.post("/api/admin/users", json={
        "student_id": "2023123", "name": "x" * 51,
    }).status_code == 400
    assert client.post("/api/admin/users", json={"student_id": "2023123"}).status_code == 201
    users = client.get("/api/admin/users").get_json()["users"]
    uid = next(u["id"] for u in users if u["student_id"] == "2023123")
    assert client.put("/api/admin/users/999999", json={
        "name": "x", "role": "user",
    }).status_code == 404
    assert client.put(f"/api/admin/users/{uid}", json={
        "name": "", "role": "user",
    }).status_code == 400
    assert client.put(f"/api/admin/users/{uid}", json={
        "name": "x", "role": "boss",
    }).status_code == 400
    assert client.delete("/api/admin/users/999999").status_code == 404


def test_admin_records_user_id_filter(client):
    make_user(client, "2023123")
    make_user(client, "2023456")
    login(client, "2023123", "2023123")
    create_record(client, product_name="甲的记录")
    login(client, "2023456", "2023456")
    create_record(client, product_name="乙的记录")

    login(client, ADMIN_ID, ADMIN_PASSWORD)
    users = client.get("/api/admin/users").get_json()["users"]
    uid = next(u["id"] for u in users if u["student_id"] == "2023123")
    records = client.get(f"/api/admin/records?user_id={uid}").get_json()["records"]
    assert [r["product_name"] for r in records] == ["甲的记录"]


OLD_INVOICES_SQL = """
CREATE TABLE invoices (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    record_id   INTEGER NOT NULL REFERENCES records(id) ON DELETE CASCADE,
    orig_name   TEXT    NOT NULL,
    stored_name TEXT    NOT NULL UNIQUE,
    uploaded_at TEXT    NOT NULL DEFAULT (datetime('now', 'localtime'))
);
CREATE INDEX idx_invoices_record ON invoices(record_id);
"""


def test_migrate_old_invoices_table(client):
    """旧 invoices 表（缺 category/remark）应被 init_db 自动加列，新字段可用。"""
    conn = db.connect()
    conn.executescript("DROP TABLE invoices;" + OLD_INVOICES_SQL)
    conn.execute(
        "INSERT INTO records (user_id, product_name, channel, amount, paid_at, payer)"
        " VALUES (1, 'x', 'jd', 1, '2026-01-01', 'self')"
    )
    conn.commit()
    conn.close()

    db.init_db()

    conn = db.connect()
    conn.execute(
        "INSERT INTO invoices (record_id, orig_name, stored_name, category, remark)"
        " VALUES (1, 'a.pdf', 'x.pdf', 'attachment', '文件备注')"
    )
    conn.commit()
    row = conn.execute(
        "SELECT category, remark FROM invoices WHERE stored_name='x.pdf'"
    ).fetchone()
    assert row[0] == "attachment"
    assert row[1] == "文件备注"
    conn.close()


def test_all_pages_render(client):
    make_user(client)
    login(client, "2023123", "2023123")
    pages = {
        "/": "login-form",
        "/user": "records-filter",
        "/reimburse": "status-filter",
        "/admin": "user-form",
        "/logs": "action-filter",
    }
    for path, marker in pages.items():
        resp = client.get(path)
        assert resp.status_code == 200, path
        assert marker in resp.get_data(as_text=True), path


def test_pages_render_with_prefix(client, monkeypatch):
    monkeypatch.setenv("FAPIAO_PREFIX", "/fapiao")
    from app import create_app

    app = create_app()
    app.config["TESTING"] = True
    c = app.test_client()
    for path in ("/fapiao/", "/fapiao/user", "/fapiao/reimburse", "/fapiao/admin", "/fapiao/logs"):
        assert c.get(path).status_code == 200, path


def test_logs_limit_1000(client):
    login(client, ADMIN_ID, ADMIN_PASSWORD)
    log_path = db.DATA_DIR / "admin.log"
    with log_path.open("a", encoding="utf-8") as fh:
        for i in range(1100):
            fh.write(json.dumps({"n": i}, ensure_ascii=False) + "\n")

    entries = client.get("/api/admin/logs").get_json()["entries"]
    assert len(entries) == 1000
    assert entries[0]["n"] == 1099
